"""Извлечь все строки из всех листов PEGAS xlsx в единый JSON для анализа."""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl

KNS2_DIR = (
    Path(r"C:\Users\User\Desktop\kns-calculator-repo")
    / "00_research"
    / "kns2_dataset_2026-05-08"
)


def find_pegas_xlsx() -> Path | None:
    for path in KNS2_DIR.rglob("*.xlsx"):
        size_mb = path.stat().st_size / 1024 / 1024
        if 6.5 < size_mb < 8.0:
            return path
    return None


def main() -> int:
    pegas_file = find_pegas_xlsx()
    if not pegas_file:
        return 1
    wb = openpyxl.load_workbook(pegas_file, data_only=True)

    out = {"_source": pegas_file.name, "sheets": {}}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            non_empty = [c for c in row if c not in (None, "")]
            if not non_empty:
                continue
            rows.append([str(c) if c is not None else "" for c in row])
        out["sheets"][sheet_name] = rows

    out_path = Path(__file__).parent / "pegas_full_dump.json"
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=2)
    print(f"Dumped {sum(len(s) for s in out['sheets'].values())} rows to {out_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
