"""Параллельная версия yadisk_extract_text.py.

Используем ProcessPoolExecutor с 6 workers.
Кешируем уже-обработанные файлы (по path) из старого extracted_text.jsonl.
Skip-эвристики:
  - PDF > 15 МБ → лимит 5 страниц (вместо 50)
  - PDF > 50 МБ → skip полностью
  - ZIP > 50 МБ → только листинг, без extract вложенного
"""
from __future__ import annotations
import json
import os
import sys
import zipfile
import time
from pathlib import Path
from concurrent.futures import ProcessPoolExecutor, as_completed

ROOT = Path(__file__).resolve().parents[1]
YA = ROOT / "02_dataset" / "_inbox" / "yadisk_kns"
OUT = YA / "extracted_text.jsonl"
LOG = YA / "extract_parallel.log"


def log(msg: str) -> None:
    print(msg, flush=True)
    LOG.parent.mkdir(parents=True, exist_ok=True)
    with LOG.open("a", encoding="utf-8") as f:
        f.write(msg + "\n")


def extract_pdf(path: Path, max_pages: int = 50) -> tuple[str, str | None]:
    try:
        from pypdf import PdfReader
    except ImportError:
        return "", "no_pdf_lib"
    try:
        r = PdfReader(str(path))
        chunks = []
        for page in r.pages[:max_pages]:
            try:
                chunks.append(page.extract_text() or "")
            except Exception:
                pass
        return "\n".join(chunks).strip(), None
    except Exception as e:
        return "", f"pdf_err:{type(e).__name__}:{str(e)[:80]}"


def extract_docx(path: Path) -> tuple[str, str | None]:
    try:
        import docx
    except ImportError:
        return "", "no_docx_lib"
    try:
        d = docx.Document(str(path))
        text = "\n".join(p.text for p in d.paragraphs if p.text.strip())
        for t in d.tables:
            for row in t.rows:
                line = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                if line:
                    text += "\n" + line
        return text.strip(), None
    except Exception as e:
        return "", f"docx_err:{type(e).__name__}:{str(e)[:80]}"


def extract_xlsx(path: Path) -> tuple[str, str | None]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return "", "no_openpyxl"
    try:
        wb = load_workbook(str(path), read_only=True, data_only=True)
        chunks = []
        for sn in wb.sheetnames[:10]:
            ws = wb[sn]
            chunks.append(f"=== {sn} ===")
            for row in ws.iter_rows(max_row=200, values_only=True):
                line = " | ".join(str(c) for c in row if c is not None)
                if line.strip():
                    chunks.append(line)
        wb.close()
        return "\n".join(chunks).strip(), None
    except Exception as e:
        return "", f"xlsx_err:{type(e).__name__}:{str(e)[:80]}"


def extract_zip_full(path: Path, max_size: int = 50 * 1024 * 1024) -> tuple[str, str | None, list]:
    """Если ZIP большой — только листинг, иначе текст из вложенных."""
    try:
        with zipfile.ZipFile(path) as zf:
            names = zf.namelist()
            chunks = [f"ZIP CONTAINS {len(names)} FILES:"]
            for n in names[:80]:
                try:
                    info = zf.getinfo(n)
                    chunks.append(f"  {n} ({info.file_size} B)")
                except Exception:
                    chunks.append(f"  {n}")
            zip_size = path.stat().st_size
            if zip_size > max_size:
                chunks.append(f"\n[zip too large {zip_size//1024}KB, listing only]")
                return "\n".join(chunks).strip(), None, names

            import tempfile
            for n in names[:20]:
                ext = Path(n).suffix.lower()
                if ext not in (".pdf", ".docx", ".xlsx"):
                    continue
                try:
                    info = zf.getinfo(n)
                    if info.file_size > 10 * 1024 * 1024:
                        continue
                    data = zf.read(n)
                    with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tf:
                        tf.write(data)
                        tp = Path(tf.name)
                    try:
                        if ext == ".pdf":
                            t, _ = extract_pdf(tp, max_pages=10)
                        elif ext == ".docx":
                            t, _ = extract_docx(tp)
                        else:
                            t, _ = extract_xlsx(tp)
                        if t:
                            chunks.append(f"\n--- {n} ---\n{t[:3000]}")
                    finally:
                        try:
                            tp.unlink()
                        except OSError:
                            pass
                except Exception:
                    pass
            return "\n".join(chunks).strip(), None, names
    except Exception as e:
        return "", f"zip_err:{type(e).__name__}:{str(e)[:80]}", []


def process_one(path_str: str) -> dict:
    p = Path(path_str)
    ext = p.suffix.lower()
    size = p.stat().st_size
    entry = {
        "path": str(p.relative_to(YA)),
        "name": p.name,
        "ext": ext,
        "size": size,
        "text": "",
        "error": None,
    }
    try:
        if ext == ".pdf":
            if size > 50 * 1024 * 1024:
                entry["error"] = "pdf_too_large_skip"
            elif size > 15 * 1024 * 1024:
                entry["text"], entry["error"] = extract_pdf(p, max_pages=5)
            else:
                entry["text"], entry["error"] = extract_pdf(p, max_pages=50)
        elif ext == ".docx":
            entry["text"], entry["error"] = extract_docx(p)
        elif ext == ".doc":
            entry["error"] = "doc_legacy_format_skip"
        elif ext == ".xlsx":
            entry["text"], entry["error"] = extract_xlsx(p)
        elif ext == ".xls":
            entry["error"] = "xls_legacy_format_skip"
        elif ext == ".zip":
            entry["text"], entry["error"], entry["zip_contents"] = extract_zip_full(p)
        elif ext in (".rar", ".7z"):
            entry["error"] = f"{ext[1:]}_archive_skip_no_lib"
        elif ext in (".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp"):
            entry["error"] = "image_no_ocr"
        elif ext in (".dwg", ".dxf"):
            entry["error"] = "cad_skip"
        else:
            entry["error"] = f"unsupported_ext:{ext}"
    except Exception as e:
        entry["error"] = f"exception:{type(e).__name__}:{str(e)[:120]}"

    if entry["text"] and len(entry["text"]) > 50000:
        entry["text"] = entry["text"][:50000] + "...[TRUNCATED]"
    return entry


def main() -> int:
    if not YA.exists():
        log(f"[ERR] {YA} not found")
        return 2
    LOG.unlink(missing_ok=True)
    OUT.parent.mkdir(parents=True, exist_ok=True)

    # Сбор всех файлов
    files = []
    for sub in ["кнс", "кнс 2", "кнс 3"]:
        d = YA / sub
        if d.exists():
            files.extend(sorted(f for f in d.iterdir() if f.is_file()))
    log(f"=== Старт PARALLEL: {len(files)} файлов, 6 workers ===")

    # Кеш уже-обработанных из старого jsonl
    cached = {}
    if OUT.exists():
        for line in OUT.read_text(encoding="utf-8", errors="replace").splitlines():
            line = line.strip()
            if not line:
                continue
            try:
                e = json.loads(line)
                cached[e.get("path", "")] = e
            except Exception:
                pass
        log(f"  cache: {len(cached)} previously processed")

    # Файлы которые НЕ в кеше
    todo = []
    for p in files:
        rel = str(p.relative_to(YA))
        if rel in cached and cached[rel].get("text") is not None:
            continue
        todo.append(str(p))

    log(f"  to process: {len(todo)} files (skipping {len(files) - len(todo)} cached)")

    # Перепишем OUT целиком: cached + new
    started = time.time()
    with OUT.open("w", encoding="utf-8") as fout:
        # cached
        for v in cached.values():
            fout.write(json.dumps(v, ensure_ascii=False) + "\n")
        # parallel
        completed = 0
        with ProcessPoolExecutor(max_workers=6) as ex:
            futures = {ex.submit(process_one, p): p for p in todo}
            for fut in as_completed(futures):
                try:
                    e = fut.result(timeout=300)
                except Exception as exc:
                    e = {
                        "path": str(Path(futures[fut]).relative_to(YA)),
                        "name": Path(futures[fut]).name,
                        "error": f"future_exc:{type(exc).__name__}",
                        "text": "",
                    }
                fout.write(json.dumps(e, ensure_ascii=False) + "\n")
                fout.flush()
                completed += 1
                if completed % 20 == 0:
                    elapsed = time.time() - started
                    rate = completed / max(elapsed, 1)
                    eta = (len(todo) - completed) / max(rate, 0.01)
                    log(f"  [{completed}/{len(todo)}] rate={rate:.1f}/s eta={eta/60:.1f}min "
                        f"{e.get('name','')[:40]} text={len(e.get('text','') or '')}")

    log(f"=== Готово за {time.time()-started:.0f}s, всего записей: {len(cached) + completed} ===")
    return 0


if __name__ == "__main__":
    sys.exit(main())
